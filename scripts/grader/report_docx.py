from __future__ import annotations

from datetime import datetime

from docx import Document
from docx.shared import Pt

from grader.action_plan import generate_personal_action_plan
from grader.models import GradingOutcome


def _add_heading(document: Document, text: str) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def write_report_docx(
    source_docx_path,
    output_docx_path,
    outcome: GradingOutcome,
    student_name: str,
    teacher_note: str | None,
) -> None:
    document = Document(str(source_docx_path))

    document.add_page_break()
    _add_heading(document, "Korrekturhilfe")

    meta = document.add_paragraph()
    meta.add_run("Datum: ").bold = True
    meta.add_run(datetime.now().strftime("%Y-%m-%d %H:%M"))
    meta.add_run("\nSchueler/in: ").bold = True
    meta.add_run(student_name)
    meta.add_run("\nProfil: ").bold = True
    meta.add_run(outcome.profile.profile_name)

    score = document.add_paragraph()
    score.add_run("Erreichte Punkte: ").bold = True
    score.add_run(f"{outcome.total_points:.2f} / {outcome.max_points:.2f}")
    score.add_run("\nNote (linear): ").bold = True
    score.add_run(str(outcome.grade))

    _add_heading(document, "Regelbasiertes Korrekturhilfe-Feedback")
    for result in outcome.results:
        status = "OK" if result.passed else "FEHLT"
        p = document.add_paragraph()
        p.add_run(f"[{status}] {result.rule.title} ").bold = True
        p.add_run(f"({result.achieved_points:.2f}/{result.rule.points:.2f} Punkte)")

        if result.rule.description:
            info = document.add_paragraph()
            info.paragraph_format.left_indent = Pt(16)
            info.add_run("Kriterium: ").bold = True
            info.add_run(result.rule.description)

        note = document.add_paragraph()
        note.paragraph_format.left_indent = Pt(16)
        note.add_run("Anmerkung: ").bold = True
        note.add_run(result.note)

    action_plan = generate_personal_action_plan(outcome, student_name)

    _add_heading(document, "Handlungsempfehlung (Juni-Abgabe)")
    section = document.add_paragraph()
    section.add_run("Fokus").bold = True
    for todo in action_plan.focus_todos:
        p = document.add_paragraph()
        p.add_run("[ ] ").bold = True
        p.add_run(todo)

    section = document.add_paragraph()
    section.add_run("2 Individuelle Erweiterungen").bold = True
    for idx, todo in enumerate(action_plan.extension_todos, start=1):
        p = document.add_paragraph()
        p.add_run("[ ] ").bold = True
        p.add_run(f"Erweiterung {idx}: {todo}")

    section = document.add_paragraph()
    section.add_run("Marschplan bis Anfang Juni").bold = True
    for todo in action_plan.timeline_todos:
        p = document.add_paragraph()
        p.add_run("[ ] ").bold = True
        p.add_run(todo)

    if teacher_note:
        _add_heading(document, "Lehrkraft-Notiz")
        document.add_paragraph(teacher_note)

    document.save(str(output_docx_path))
